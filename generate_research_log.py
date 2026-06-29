"""Generate research process Word document for scheduling_research project."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.bold = bold


def add_code(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(10)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val


def build_document() -> Document:
    doc = Document()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("动态柔性作业车间调度（DFJSP）\n大模型微调研究过程记录")
    run.bold = True
    run.font.size = Pt(16)

    add_para(doc, f"项目名称：scheduling_research")
    add_para(doc, f"记录日期：{date.today().isoformat()}")
    add_para(doc, "硬件环境：NVIDIA GeForce RTX 3050 Laptop GPU（4 GB 显存）")
    add_para(doc, "软件环境：Windows 10/11，Python 3.12，CUDA 11.8（驱动支持 12.7）")
    doc.add_paragraph()

    add_heading(doc, "一、研究背景与目标", 1)
    add_para(
        doc,
        "本课题面向动态柔性作业车间调度问题（Dynamic Flexible Job Shop Scheduling, DFJSP）。"
        "场景为：4 台机器、6 个初始作业，每作业 3–5 道工序，加工时间 1–15 随机；"
        "在调度进度约 50% 时插入 1–2 个急单，需重新优化未完成工序与急单的调度方案。"
    )
    add_para(
        doc,
        "技术路线：使用 OR-Tools CP-SAT 生成近似最优调度标签 → 构造文本描述 + JSON 最优动作 → "
        "在 RTX 3050（4 GB 显存）上对 unsloth/Qwen2.5-1.5B 进行 4-bit QLoRA 微调，"
        "使小模型学习“读问题描述 → 输出最优调度 JSON”。"
    )

    add_heading(doc, "二、项目结构与已完成工作", 1)
    add_table(
        doc,
        ["文件", "说明"],
        [
            ["requirements.txt / install.ps1", "分步安装依赖，防止 unsloth 覆盖 PyTorch"],
            ["check_env.py", "检查 CUDA、显存、PyTorch 版本"],
            ["data_generator.py", "生成 200 条训练 + 50 条测试 JSONL 数据"],
            ["train_unsloth.py", "4-bit LoRA 微调脚本（max_seq_length=512）"],
            ["run_train.ps1", "一键训练（含 HF 镜像与 UTF-8 设置）"],
            ["train_data.jsonl / test_data.jsonl", "OR-Tools 标注的调度数据"],
            ["lora_output/", "训练完成后保存的 LoRA 适配器"],
        ],
    )
    doc.add_paragraph()

    add_heading(doc, "三、研究过程记录", 1)

    add_heading(doc, "阶段 1：环境搭建", 2)
    add_para(doc, "目标：在 4 GB 显存笔记本上搭建可用的 PyTorch + Unsloth 环境。")
    add_para(doc, "操作步骤：", bold=True)
    add_code(
        doc,
        "cd D:\\work\\scheduling_research\n"
        "python -m venv .venv\n"
        ".\\.venv\\Scripts\\Activate.ps1\n"
        ".\\install.ps1\n"
        "python check_env.py"
    )
    add_para(doc, "预期结果：CUDA available: True，PyTorch 2.4.1+cu118，显存 4 GB 识别正常。")

    add_heading(doc, "阶段 2：数据生成", 2)
    add_para(doc, "使用 OR-Tools 对动态 FJSP 实例求解，生成监督微调数据。")
    add_code(doc, "python data_generator.py\npython preview_data.py")
    add_para(
        doc,
        "数据格式（JSONL 每行）：instance_id、description（中文问题描述）、"
        "optimal_action（含 makespan 与 schedule 序列）。"
        "共 200 条训练、50 条测试；单条约 30 道工序，求解限时 3 秒/实例。"
    )

    add_heading(doc, "阶段 3：LoRA 微调", 2)
    add_para(doc, "关键超参（受 4 GB 显存约束）：", bold=True)
    add_table(
        doc,
        ["参数", "取值", "说明"],
        [
            ["模型", "unsloth/Qwen2.5-1.5B", "4-bit 量化加载"],
            ["max_seq_length", "512", "不超过 512"],
            ["LoRA r", "8", "低秩适配"],
            ["batch × accum", "1 × 8", "等效 batch=8"],
            ["gradient_checkpointing", "True", "节省显存"],
            ["fp16", "True", "混合精度"],
            ["max_steps", "200", "先跑通流程"],
        ],
    )
    doc.add_paragraph()
    add_code(doc, ".\\run_train.ps1")
    add_para(
        doc,
        "冒烟测试（5 step）结果：loss≈2.08，峰值显存≈1.95 GB / 4 GB，"
        "LoRA 成功保存至 ./lora_output。"
    )

    add_heading(doc, "四、遇到的问题与解决方案", 1)

    problems = [
        (
            "问题 1：check_env 显示 CUDA available: False（torch 2.10.0+cpu）",
            "原因",
            "pip install -r requirements.txt 时，unsloth 从 PyPI 拉取了 CPU 版 PyTorch，"
            "覆盖了已安装的 torch 2.4.1+cu118。",
            "解决",
            "分步安装：先装 CUDA 版 PyTorch，再 pip install --no-deps 装 unsloth。\n"
            "重装命令：\n"
            "pip uninstall torch torchvision torchaudio xformers -y\n"
            "pip install torch==2.4.1 torchvision==0.19.1 torchaudio==2.4.1 "
            "--index-url https://download.pytorch.org/whl/cu118\n"
            "pip install xformers==0.0.27.post2 --index-url https://download.pytorch.org/whl/cu118 --no-deps",
        ),
        (
            "问题 2：OSError 加载 fbgemm.dll 失败（WinError 126）",
            "原因",
            "xformers 安装时试图重装 torch 2.4.0，导致 torch 文件混装、DLL 依赖损坏。",
            "解决",
            "完整卸载后重装 torch 2.4.1+cu118 全家桶；xformers 必须加 --no-deps。",
        ),
        (
            "问题 3：train_unsloth.py 报错 torch.int1 / transformers 5.5",
            "原因",
            "unsloth 2026.x 自动升级了 transformers 5.5.0 和 torchao 0.17.0，"
            "二者需要 torch≥2.11，与当前 torch 2.4.1 不兼容。",
            "解决",
            "卸载 torchao；降级 transformers==4.51.3、trl==0.15.2；"
            "锁定 unsloth==2025.6.8（勿用 2026.x）。",
        ),
        (
            "问题 4：unsloth 2026.x 报 torch._inductor.config 不存在",
            "原因",
            "最新 unsloth_zoo 依赖 PyTorch 2.5+ 的 inductor 接口，Windows 上 torch 2.4.1 无此属性。",
            "解决",
            "降级为 unsloth==2025.6.8 + unsloth_zoo==2025.6.8（与 torch 2.4.1 兼容）。",
        ),
        (
            "问题 5：HuggingFace 连接 huggingface.co 超时",
            "原因",
            "国内网络访问 HuggingFace 官方站不稳定。",
            "解决",
            "训练前设置镜像：\n"
            "$env:HF_ENDPOINT = \"https://hf-mirror.com\"\n"
            "（已写入 run_train.ps1）",
        ),
        (
            "问题 6：Windows GBK 解码错误 / UnslothSFTTrainer 多进程 tokenize 失败",
            "原因",
            "① unsloth 读取 trl 源码时用 GBK 解码失败；"
            "② datasets 多进程 map 在 Windows 子进程找不到 UnslothSFTTrainer 模块。",
            "解决",
            "① 设置 $env:PYTHONUTF8=\"1\"；\n"
            "② train_unsloth.py 中改为单进程预 tokenize，"
            "并设置 dataset_kwargs={\"skip_prepare_dataset\": True}。",
        ),
        (
            "问题 7：accelerate 版本冲突（FP8BackendType / data_seed）",
            "原因",
            "accelerate 1.14.0 与 transformers 4.51 补丁不兼容；"
            "0.34.2 又缺少 data_seed 支持。",
            "解决",
            "固定 accelerate==1.1.1，使用 pip install --no-deps 安装，避免覆盖 torch。",
        ),
    ]

    for title, lbl1, reason, lbl2, solution in problems:
        add_heading(doc, title, 2)
        add_para(doc, lbl1 + "：", bold=True)
        add_para(doc, reason)
        add_para(doc, lbl2 + "：", bold=True)
        add_code(doc, solution)

    add_heading(doc, "五、最终可用依赖版本（锁定）", 1)
    add_table(
        doc,
        ["包名", "版本", "备注"],
        [
            ["torch / torchvision / torchaudio", "2.4.1+cu118", "PyTorch 官方 cu118 源"],
            ["transformers", "4.51.3", "与 unsloth 2025.6.8 兼容"],
            ["trl", "0.15.2", "SFT 训练"],
            ["accelerate", "1.1.1", "--no-deps 安装"],
            ["unsloth / unsloth_zoo", "2025.6.8", "勿用 2026.x"],
            ["xformers", "0.0.27.post2+cu118", "--no-deps 安装"],
            ["bitsandbytes", "≥0.45.5", "4-bit 量化"],
        ],
    )
    doc.add_paragraph()

    add_heading(doc, "六、常用命令速查", 1)
    add_code(
        doc,
        "# 环境检查\n"
        "python check_env.py\n\n"
        "# 生成数据\n"
        "python data_generator.py\n\n"
        "# 训练（推荐）\n"
        ".\\run_train.ps1\n\n"
        "# 手动训练\n"
        "$env:PYTHONUTF8 = \"1\"\n"
        "$env:HF_ENDPOINT = \"https://hf-mirror.com\"\n"
        "python train_unsloth.py"
    )

    add_heading(doc, "七、后续计划", 1)
    add_para(doc, "1. 完成 200 step 全量训练，观察 loss 曲线与显存峰值。")
    add_para(doc, "2. 在 test_data.jsonl 上评估模型输出 JSON 的 makespan 与 OR-Tools 标签的差距。")
    add_para(doc, "3. 尝试缩短 description 或压缩 schedule 表示，提升 512 token 内的有效信息密度。")
    add_para(doc, "4. 若效果不足，可考虑增大数据量或引入 prompt-completion 分离 loss。")

    add_heading(doc, "八、参考资料", 1)
    add_para(doc, "• Unsloth 官方安装文档：https://unsloth.ai/docs/get-started/install/pip-install")
    add_para(doc, "• OR-Tools CP-SAT 作业车间示例")
    add_para(doc, "• 项目代码目录：D:\\work\\scheduling_research")
    add_para(doc, "• 内部参考文档：adapt pruner.docx、C-PRUNE.docx（研究过程记录格式）")

    return doc


def main() -> None:
    out = Path(__file__).parent / "research_log_DFJSP.docx"
    doc = build_document()
    doc.save(out)
    print(f"已生成: {out}")


if __name__ == "__main__":
    main()
